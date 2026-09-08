"""Unit tests for document parsing and semantic career-unit chunking."""

import io
from pathlib import Path
import docx
import pypdf
import pytest

from src.data_ingestion.document_parser import DocumentParser
from src.data_ingestion.semantic_chunker import CareerChunker
from src.models.profile import CareerChunk


SAMPLE_RESUME_TEXT = """
Alex Mercer
alex.mercer@example.com | (555) 019-2834 | Washington, DC

PROFESSIONAL SUMMARY
Data Scientist with 3+ years of experience in natural language processing, machine learning system development, and data pipelines.

WORK EXPERIENCE
Senior Machine Learning Engineer | Federal Analytics Corp | Jan 2022 - Present
- Built transformer-based BERT pipelines for federal document classification using PyTorch and Hugging Face.
- Designed distributed data pipelines on AWS with PostgreSQL and Docker.
- Reduced model inference latency by 45% using ONNX runtime and quantization.

Junior Data Analyst | Tech Defense Solutions | Jun 2020 - Dec 2021
- Developed SQL reporting dashboards and automated ETL workflows in Python.
- Trained scikit-learn models for predictive equipment maintenance.

PROJECTS
OpenSource NLP Tool | 2023
- Implemented retrieval-augmented generation (RAG) using LangChain, ChromaDB, and Python.
- Evaluated cosine similarity and BM25 reranking across 50,000 legal contracts.

EDUCATION
Master of Science in Computer Science | George Washington University | 2020 - 2022
- Focus on Machine Learning and Natural Language Processing. GPA: 3.9/4.0.

Bachelor of Science in Information Systems | University of Maryland | 2016 - 2020

TECHNICAL SKILLS
- Languages: Python, SQL, R, Bash
- ML/AI: PyTorch, scikit-learn, Hugging Face Transformers, BERT, NLP
- Cloud & DevOps: AWS, Docker, Git, CI/CD, Linux

CERTIFICATIONS
- AWS Certified Machine Learning - Specialty (2023)
"""


def test_clean_text():
    dirty = "Line 1   with   spaces\r\n\r\n\r\n\r\nLine 2 with \u00a0 special chars."
    cleaned = DocumentParser.clean_text(dirty)
    assert "Line 1 with spaces" in cleaned
    assert "Line 2 with special chars." in cleaned
    assert "\r" not in cleaned


def test_txt_parsing():
    parsed = DocumentParser.parse_document("resume.txt", SAMPLE_RESUME_TEXT.encode("utf-8"))
    assert "Alex Mercer" in parsed
    assert "PyTorch" in parsed


def test_docx_parsing():
    doc = docx.Document()
    doc.add_heading("Alex Mercer", level=1)
    doc.add_paragraph("Master of Science in Computer Science")
    doc.add_paragraph("Skills: Python, SQL, PyTorch")

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    parsed = DocumentParser.parse_document("test_resume.docx", docx_bytes)
    assert "Alex Mercer" in parsed
    assert "Computer Science" in parsed
    assert "PyTorch" in parsed


def test_semantic_chunking_units():
    chunks = CareerChunker.chunk_document(SAMPLE_RESUME_TEXT, "resume.txt")
    assert len(chunks) >= 5, f"Expected at least 5 career chunks, got {len(chunks)}"

    # Verify sections exist
    sections = {c.section for c in chunks}
    assert "work_experience" in sections
    assert "education" in sections
    assert "skills" in sections
    assert "projects" in sections

    # Check work experience chunk metadata
    work_chunks = [c for c in chunks if c.section == "work_experience"]
    assert len(work_chunks) >= 2
    titles = [c.job_title for c in work_chunks if c.job_title]
    assert any("Senior Machine Learning Engineer" in t for t in titles)

    # Check skills extracted
    all_skills = set()
    for c in chunks:
        all_skills.update(c.skills)
    assert "Python" in all_skills
    assert "Pytorch" in all_skills or "PYTORCH" in all_skills
    assert "Sql" in all_skills or "SQL" in all_skills


def test_citation_formatting():
    chunk = CareerChunk(
        chunk_id="resume.txt_001",
        source_document="resume.txt",
        section="work_experience",
        job_title="ML Engineer",
        organization="Federal Corp",
        date="2022 - Present",
        skills=["Python", "PyTorch"],
        content="Engineered machine learning pipelines.",
    )
    citation = chunk.to_citation()
    assert "resume.txt" in citation
    assert "Work Experience" in citation
    assert "ML Engineer" in citation
    assert "resume.txt_001" in citation
